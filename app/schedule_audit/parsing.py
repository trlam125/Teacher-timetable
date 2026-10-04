from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


MAX_ARCHIVE_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".docx"}


class ScheduleAuditParseError(ValueError):
    pass


@dataclass
class RawLesson:
    day_text: str
    session_text: str
    period_text: str
    class_text: str
    lesson_text: str = ""
    subject_text: str = ""
    teacher_text: str = ""
    room_text: str = ""
    source: str = ""
    origin: str = "aggregate"


def normalize_text(value: Any) -> str:
    """Normalize for fuzzy/search matching; Vietnamese accents are intentionally ignored."""
    text = str(value or "").strip().casefold().replace("đ", "d")
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _identity_text(value: Any) -> str:
    """Normalize an entity identity without throwing away Vietnamese diacritics."""
    text = unicodedata.normalize("NFC", str(value or "").strip().casefold())
    text = "".join(ch if ch.isalnum() else " " for ch in text)
    return re.sub(r"\s+", " ", text).strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _archive_size_guard(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            total = sum(max(0, item.file_size) for item in zf.infolist())
    except zipfile.BadZipFile as exc:
        raise ScheduleAuditParseError("File nén không hợp lệ hoặc đã bị hỏng.") from exc
    if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
        raise ScheduleAuditParseError(
            "File có dữ liệu giải nén quá lớn để kiểm tra an toàn."
        )


def _read_excel_tables(
    content: bytes, suffix: str
) -> list[tuple[str, list[list[str]]]]:
    if suffix in {".xlsx", ".xlsm"}:
        _archive_size_guard(content)
        try:
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
        except Exception as exc:
            raise ScheduleAuditParseError(
                "Không đọc được file Excel. Hãy kiểm tra file có bị hỏng không."
            ) from exc
        tables: list[tuple[str, list[list[str]]]] = []
        try:
            for sheet in workbook.worksheets:
                rows = [
                    [_cell_text(value) for value in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                tables.append((sheet.title, rows))
        finally:
            workbook.close()
        return tables

    try:
        import xlrd  # type: ignore
    except ImportError as exc:
        raise ScheduleAuditParseError(
            "File .xls cần thư viện xlrd. Hãy chạy lại pip install -r requirements.txt hoặc lưu file thành .xlsx."
        ) from exc
    try:
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as exc:
        raise ScheduleAuditParseError("Không đọc được file Excel .xls.") from exc
    tables = []
    try:
        for sheet in book.sheets():
            rows = [
                [_cell_text(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
                for r in range(sheet.nrows)
            ]
            tables.append((sheet.name, rows))
    finally:
        book.release_resources()
    return tables


def _decode_delimited(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1258", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ScheduleAuditParseError("Không nhận diện được bảng mã của file CSV/TSV.")


def _read_delimited_table(
    content: bytes, suffix: str
) -> list[tuple[str, list[list[str]]]]:
    text = _decode_delimited(content)
    sample = text[:8192]
    delimiter = "\t" if suffix == ".tsv" else ","
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        pass
    rows = [
        [_cell_text(value) for value in row]
        for row in csv.reader(io.StringIO(text), delimiter=delimiter)
    ]
    return [("Dữ liệu", rows)]


def _docx_text(element: ET.Element, namespace: dict[str, str]) -> str:
    paragraph_tag = f"{{{namespace['w']}}}p"
    paragraphs = (
        [element]
        if element.tag == paragraph_tag
        else element.findall(".//w:p", namespace)
    )
    if paragraphs:
        chunks = [
            "".join(
                node.text or "" for node in paragraph.findall(".//w:t", namespace)
            ).strip()
            for paragraph in paragraphs
        ]
        return re.sub(r"\s+", " ", " ".join(chunk for chunk in chunks if chunk)).strip()
    texts = [node.text or "" for node in element.findall(".//w:t", namespace)]
    return re.sub(r"\s+", " ", "".join(texts)).strip()


def _read_docx_tables(content: bytes) -> list[tuple[str, list[list[str]]]]:
    _archive_size_guard(content)
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            xml = zf.read("word/document.xml")
    except (KeyError, zipfile.BadZipFile) as exc:
        raise ScheduleAuditParseError("Không đọc được tài liệu Word .docx.") from exc
    try:
        root = ET.fromstring(xml)
    except ET.ParseError as exc:
        raise ScheduleAuditParseError(
            "Nội dung XML của tài liệu Word không hợp lệ."
        ) from exc

    body = root.find("w:body", namespace)
    if body is None:
        return []
    tables: list[tuple[str, list[list[str]]]] = []
    last_paragraph = ""
    table_no = 0
    paragraph_tag = f"{{{namespace['w']}}}p"
    table_tag = f"{{{namespace['w']}}}tbl"
    for child in body:
        if child.tag == paragraph_tag:
            text = _docx_text(child, namespace)
            if text:
                last_paragraph = text
            continue
        if child.tag != table_tag:
            continue
        table_no += 1
        rows: list[list[str]] = []
        for tr in child.findall("w:tr", namespace):
            row = [_docx_text(tc, namespace) for tc in tr.findall("w:tc", namespace)]
            rows.append(row)
        title = last_paragraph or f"Bảng {table_no}"
        tables.append((title, rows))
    return tables


def read_tables(
    filename: str, content: bytes
) -> tuple[str, list[tuple[str, list[list[str]]]]]:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        if suffix == ".doc":
            raise ScheduleAuditParseError(
                "Word .doc đời cũ chưa được hỗ trợ. Hãy lưu lại dưới dạng .docx."
            )
        raise ScheduleAuditParseError(
            "Định dạng chưa hỗ trợ. Dùng .xlsx, .xlsm, .xls, .csv, .tsv hoặc .docx."
        )
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return suffix.lstrip(".").upper(), _read_excel_tables(content, suffix)
    if suffix in {".csv", ".tsv"}:
        return suffix.lstrip(".").upper(), _read_delimited_table(content, suffix)
    return "DOCX", _read_docx_tables(content)


def _header_kind(value: str) -> str:
    text = normalize_text(value)
    if text in {"thu", "ngay", "day", "weekday"} or text.startswith("thu "):
        return "day"
    if text in {"buoi", "ca", "session", "shift"}:
        return "session"
    if text in {"tiet", "tiet hoc", "period", "lesson", "lesson no"}:
        return "period"
    if text in {"lop", "lop hoc", "class", "classroom"}:
        return "class"
    if text in {"mon", "mon hoc", "subject"}:
        return "subject"
    if text in {"giao vien", "gv", "teacher"}:
        return "teacher"
    if text in {"phong", "phong hoc", "room"}:
        return "room"
    if text in {"noi dung", "lich hoc", "bai hoc", "lesson info", "noi dung tiet"}:
        return "lesson"
    return ""


def _day_index(value: str) -> int | None:
    text = normalize_text(value)
    if not text:
        return None
    if text in {"cn", "chu nhat", "sunday", "sun"}:
        return 6
    english = {
        "monday": 0,
        "mon": 0,
        "tuesday": 1,
        "tue": 1,
        "wednesday": 2,
        "wed": 2,
        "thursday": 3,
        "thu": 3,
        "friday": 4,
        "fri": 4,
        "saturday": 5,
        "sat": 5,
    }
    if text in english:
        return english[text]
    match = re.search(r"(?:thu\s*)?([2-8])\b", text)
    if match:
        number = int(match.group(1))
        return 6 if number == 8 else number - 2
    if text.isdigit():
        number = int(text)
        if 2 <= number <= 8:
            return 6 if number == 8 else number - 2
    return None


def _session_index(value: str, session_count: int) -> int | None:
    if session_count <= 1:
        return 0
    text = normalize_text(value)
    if not text:
        return None
    # File TKB thuc te thuong ghi "Sang 2" / "Chieu 2", trong do so 2 la thu
    # chu khong phai so buoi. Uu tien tu Sang/Chieu truoc khi doc bat ky chu so nao.
    if text == "s" or text.startswith("sang") or "morning" in text:
        return 0
    if text == "c" or text.startswith("chieu") or "afternoon" in text:
        return 1 if session_count > 1 else 0
    if any(token in text for token in ("buoi 1", "ca 1")):
        return 0
    if any(token in text for token in ("buoi 2", "ca 2")):
        return 1 if session_count > 1 else 0
    match = re.search(r"(?:buoi|ca|session)\s*([1-9])\b", text)
    if match:
        index = int(match.group(1)) - 1
        return index if 0 <= index < session_count else None
    return None


def _period_number(value: str) -> int | None:
    text = normalize_text(value)
    if not text:
        return None
    match = re.search(r"\b(\d{1,2})\b", text)
    return int(match.group(1)) if match else None


def _implicit_session_hints(
    rows: list[list[str]],
    *,
    start_index: int,
    period_col: int,
    session_col: int | None = None,
    day_col: int | None = None,
) -> dict[int, str]:
    """Infer Sang/Chieu when a timetable repeats period numbers without a session column.

    Grid/wide schedules commonly list periods 1..5 for the morning and then restart
    at 1..5 for the afternoon.  The reset is trustworthy only for layouts where one
    source row represents one timetable period.  Long/assignment tables are excluded
    because several classes may legitimately repeat the same period on adjacent rows.
    """
    if session_col is not None:
        explicit_sessions = {
            _session_index(_cell_text(values[session_col]), 2)
            for values in rows[start_index:]
            if session_col < len(values) and _cell_text(values[session_col])
        }
        explicit_sessions.discard(None)
        # When both sessions are explicitly represented, the parser should trust
        # the source labels. If the column is blank or only contains one carried
        # label, period resets can still supply the missing session boundary.
        if {0, 1}.issubset(explicit_sessions):
            return {}

    hints: dict[int, str] = {}
    state: dict[str, dict[str, int]] = {}
    last_day = ""
    saw_second_session = False
    for row_index in range(start_index, len(rows)):
        values = rows[row_index]
        if period_col >= len(values):
            continue
        period = _period_number(_cell_text(values[period_col]))
        if period is None:
            continue

        if day_col is None:
            day_key = "__grid__"
        else:
            day_value = _cell_text(values[day_col]) if day_col < len(values) else ""
            if day_value:
                last_day = day_value
            if not last_day:
                continue
            day_key = normalize_text(last_day) or last_day

        current = state.setdefault(day_key, {"session": 0, "last": 0, "peak": 0})
        last_period = int(current["last"] or 0)
        peak = int(current["peak"] or 0)
        # A genuine morning->afternoon restart normally follows at least period 3
        # and restarts at period 1/2. This avoids treating harmless duplicate rows
        # as a new session.
        if last_period and peak >= 3 and period <= 2 and period < last_period:
            current["session"] = min(1, int(current["session"]) + 1)
            current["peak"] = 0
        current["last"] = period
        current["peak"] = max(int(current["peak"]), period)
        if int(current["session"]) == 1:
            saw_second_session = True
        hints[row_index] = "Chieu" if int(current["session"]) == 1 else "Sang"

    # Do not invent a morning label for ordinary one-session files.  Returning no
    # hints preserves the previous behaviour unless a real period reset was seen.
    return hints if saw_second_session else {}


def _resolve_slot(
    day_text: str, session_text: str, period_text: str, project: dict[str, Any]
) -> int | None:
    days = int(project["days"])
    sessions = int(project["sessions"])
    periods = int(project["periods"])
    day = _day_index(day_text)
    period = _period_number(period_text)
    if day is None or period is None or not (0 <= day < days):
        return None
    session = _session_index(session_text, sessions)
    if session is None and 1 <= period <= sessions * periods:
        session = (period - 1) // periods
        period = ((period - 1) % periods) + 1
    if session is None:
        session = 0 if sessions == 1 else None
    if session is None or not (0 <= session < sessions) or not (1 <= period <= periods):
        return None
    return day * sessions * periods + session * periods + period - 1


def slot_label(slot: int, project: dict[str, Any]) -> str:
    periods = int(project["periods"])
    sessions = int(project["sessions"])
    per_day = periods * sessions
    day = slot // per_day
    inside = slot % per_day
    session = inside // periods
    period = inside % periods + 1
    day_name = ["Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "CN"][day]
    if sessions == 1:
        return f"{day_name}, tiết {period}"
    session_name = (
        "Sáng" if session == 0 else ("Chiều" if session == 1 else f"Buổi {session + 1}")
    )
    return f"{day_name}, {session_name.lower()}, tiết {period}"


def _trim_matrix(rows: list[list[str]]) -> list[list[str]]:
    cleaned = []
    for row in rows:
        values = list(row)
        while values and not _cell_text(values[-1]):
            values.pop()
        cleaned.append([_cell_text(value) for value in values])
    while cleaned and not any(cleaned[-1]):
        cleaned.pop()
    return cleaned


def _parse_long_table(title: str, rows: list[list[str]]) -> list[RawLesson]:
    for header_index, row in enumerate(rows[:30]):
        kinds: dict[str, int] = {}
        for col, value in enumerate(row):
            kind = _header_kind(value)
            if kind and kind not in kinds:
                kinds[kind] = col
        if not {"day", "period", "class"}.issubset(kinds):
            continue
        if not ({"subject", "lesson", "teacher"} & set(kinds)):
            continue
        parsed: list[RawLesson] = []
        last_day = ""
        last_session = ""
        for row_no, values in enumerate(
            rows[header_index + 1 :], start=header_index + 2
        ):
            get = lambda key: (
                values[kinds[key]] if key in kinds and kinds[key] < len(values) else ""
            )
            day = get("day") or last_day
            session = get("session") or last_session
            period = get("period")
            class_text = get("class")
            subject = get("subject")
            lesson = get("lesson")
            teacher = get("teacher")
            room = get("room")
            if get("day"):
                last_day = get("day")
            if get("session"):
                last_session = get("session")
            if not period or not class_text or not (subject or lesson or teacher):
                continue
            parsed.append(
                RawLesson(
                    day,
                    session,
                    period,
                    class_text,
                    lesson,
                    subject,
                    teacher,
                    room,
                    f"{title} · dòng {row_no}",
                    "aggregate",
                )
            )
        if parsed:
            return parsed
    return []


_WIDE_TABLE_AUXILIARY_HEADERS = {
    # Columns that describe the row/schedule rather than a class/teacher entity.
    "gv nghi",
    "giao vien nghi",
    "gv vang",
    "giao vien vang",
    "ghi chu",
    "chu thich",
    "note",
    "notes",
    "remark",
    "remarks",
    "phong",
    "phong hoc",
    "room",
    "rooms",
    "stt",
    "so tt",
    "so thu tu",
    "tt",
}


def _is_wide_table_auxiliary_header(value: str) -> bool:
    text = normalize_text(value)
    if not text:
        return False
    if text in _WIDE_TABLE_AUXILIARY_HEADERS:
        return True
    # Generic long-table metadata headers must never become entity columns
    # when the same sheet is interpreted as a wide timetable.
    return _header_kind(value) in {"class", "subject", "teacher", "room", "lesson"}


def _issue(
    code: str,
    severity: str,
    title: str,
    detail: str,
    *,
    slot: int | None = None,
    project: dict[str, Any] | None = None,
    source: str = "",
    entity: str = "",
) -> dict[str, Any]:
    return {
        "code": code,
        "severity": severity,
        "title": title,
        "detail": detail,
        "slot": slot,
        "slot_label": slot_label(slot, project)
        if slot is not None and project is not None
        else "",
        "source": source,
        "entity": entity,
    }


# ===== Standalone import/audit (khong phu thuoc project co san) =====

_STANDALONE_SUBJECT_HINTS = {
    "hdtnhn": "HĐTNHN",
    "tnhn": "TNHN",
    "hdtn": "HĐTN",
    "toan": "Toán",
    "ngu van": "Ngữ văn",
    "n van": "Ngữ văn",
    "van": "Ngữ văn",
    "tieng anh": "Tiếng Anh",
    "t anh": "Tiếng Anh",
    "anh": "Tiếng Anh",
    "vat ly": "Vật lý",
    "ly": "Vật lý",
    "hoa hoc": "Hóa học",
    "hoa": "Hóa học",
    "sinh hoc": "Sinh học",
    "sinh": "Sinh học",
    "lich su": "Lịch sử",
    "su": "Lịch sử",
    "dia ly": "Địa lý",
    "dia": "Địa lý",
    "lsdl": "LSĐL",
    "gdcd": "GDCD",
    "gdktpl": "GDKTPL",
    "gddp": "GDĐP",
    "tin hoc": "Tin học",
    "tin": "Tin học",
    "cong nghe": "Công nghệ",
    "c nghe": "Công nghệ",
    "cn": "Công nghệ",
    "the duc": "Thể dục",
    "gdtc": "GDTC",
    "am nhac": "Âm nhạc",
    "nt nhac": "Âm nhạc",
    "my thuat": "Mỹ thuật",
    "nt mt": "Mỹ thuật",
    "quoc phong": "Quốc phòng",
    "gdqp": "GDQP",
    "khtn": "KHTN",
    "khxh": "KHXH",
    "trai nghiem": "Trải nghiệm",
    "chao co": "Chào cờ",
    "sinh hoat": "Sinh hoạt",
}

# Cac mau o TKB pho bien: "N.Van Que", "KHTN(S) N.Tam", "HDTNHN Que"...
# Match prefix tren chuoi goc de giu lai ten giao vien phia sau chinh xac.
_STANDALONE_SUBJECT_PREFIX_PATTERNS: tuple[tuple[re.Pattern[str], str], ...] = tuple(
    (re.compile(pattern, re.IGNORECASE), subject)
    for pattern, subject in (
        # Put longer/full subject names before shorter prefixes. The trailing
        # word boundary prevents cases such as "Sinh hoạt" being consumed as
        # "Sinh" (Sinh học).
        (r"^\s*HĐTNHN\b\s*[-:]?\s*", "HĐTNHN"),
        (r"^\s*TNHN\b\s*[-:]?\s*", "TNHN"),
        (r"^\s*HĐTN\b\s*[-:]?\s*", "HĐTN"),
        (r"^\s*Sinh\s+hoạt\b\s*[-:]?\s*", "Sinh hoạt"),
        (r"^\s*Chào\s+cờ\b\s*[-:]?\s*", "Chào cờ"),
        (r"^\s*Trải\s+nghiệm\b\s*[-:]?\s*", "Trải nghiệm"),
        (r"^\s*Công\s+nghệ\b\s*[-:]?\s*", "Công nghệ"),
        (r"^\s*Thể\s+dục\b\s*[-:]?\s*", "Thể dục"),
        (r"^\s*Âm\s+nhạc\b\s*[-:]?\s*", "Âm nhạc"),
        (r"^\s*Mỹ\s+thuật\b\s*[-:]?\s*", "Mỹ thuật"),
        (r"^\s*Quốc\s+phòng\b\s*[-:]?\s*", "Quốc phòng"),
        (r"^\s*N\.?\s*Văn\b\s*[-:]?\s*", "Ngữ văn"),
        (r"^\s*(?:Ngữ\s+)?Văn\b\s*[-:]?\s*", "Ngữ văn"),
        (r"^\s*T\.?\s*Anh\b\s*[-:]?\s*", "Tiếng Anh"),
        (r"^\s*(?:Tiếng\s+)?Anh\b\s*[-:]?\s*", "Tiếng Anh"),
        (r"^\s*(?:Vật\s+)?Lý\b\s*[-:]?\s*", "Vật lý"),
        (r"^\s*Hóa(?:\s+học)?\b\s*[-:]?\s*", "Hóa học"),
        (r"^\s*Sinh(?:\s+học)?\b\s*[-:]?\s*", "Sinh học"),
        (r"^\s*(?:Lịch\s+)?Sử\b\s*[-:]?\s*", "Lịch sử"),
        (r"^\s*Địa(?:\s+lý)?\b\s*[-:]?\s*", "Địa lý"),
        (r"^\s*C\.?\s*Nghệ\b\s*[-:]?\s*", "Công nghệ"),
        (r"^\s*CN\b\s*[-:]?\s*", "Công nghệ"),
        (r"^\s*LSĐL(?:\s*\([^)]*\))?\s*[-:]?\s*", "LSĐL"),
        (r"^\s*KHTN(?:\s*\([^)]*\))?\s*[-:]?\s*", "KHTN"),
        (r"^\s*KHXH(?:\s*\([^)]*\))?\s*[-:]?\s*", "KHXH"),
        (r"^\s*NT\s*\(\s*Nhạc\s*\)\s*[-:]?\s*", "Âm nhạc"),
        (r"^\s*NT\s*\(\s*MT\s*\)\s*[-:]?\s*", "Mỹ thuật"),
        (r"^\s*GDĐP\b\s*[-:]?\s*", "GDĐP"),
        (r"^\s*GDDP\b\s*[-:]?\s*", "GDĐP"),
        (r"^\s*GDKTPL\b\s*[-:]?\s*", "GDKTPL"),
        (r"^\s*GDTC\b\s*[-:]?\s*", "GDTC"),
        (r"^\s*GDCD\b\s*[-:]?\s*", "GDCD"),
        (r"^\s*GDQP\b\s*[-:]?\s*", "GDQP"),
        (r"^\s*Tin(?:\s+học)?\b\s*[-:]?\s*", "Tin học"),
        (r"^\s*Toán\b\s*[-:]?\s*", "Toán"),
    )
)


def _standalone_subject_prefix(value: str) -> tuple[str, str]:
    text = unicodedata.normalize("NFC", _cell_text(value))
    if not text:
        return "", ""

    # First handle normalized aliases so the parser behaves the same for
    # accented and unaccented text (e.g. "Công nghệ" / "cong nghe").
    # Parenthesized forms are left to the regexes below because they know how
    # to consume qualifiers such as KHTN(S) or NT(Nhạc) without treating the
    # qualifier as part of the teacher name.
    normalized = normalize_text(text)
    for hint in sorted(_STANDALONE_SUBJECT_HINTS, key=len, reverse=True):
        if normalized != hint and not normalized.startswith(f"{hint} "):
            continue
        prefix_end = None
        for end in range(1, len(text) + 1):
            if normalize_text(text[:end]) == hint:
                prefix_end = end
                break
        if prefix_end is None:
            continue
        remainder_raw = text[prefix_end:]
        if "(" in text[:prefix_end] or remainder_raw.lstrip().startswith("("):
            continue
        return _STANDALONE_SUBJECT_HINTS[hint], remainder_raw.strip(" -–—|;/: ")

    for pattern, subject in _STANDALONE_SUBJECT_PREFIX_PATTERNS:
        match = pattern.match(text)
        if match:
            return subject, text[match.end() :].strip(" -–—|;/: ")
    return "", ""


def _standalone_clean_entity_heading(value: str) -> str:
    """Remove timetable/entity labels while preserving the original Vietnamese name."""
    original = _cell_text(value)
    text = unicodedata.normalize("NFC", original)
    timetable_label = r"(?:tkb|thời\s*khóa\s*biểu|thoi\s*khoa\s*bieu)"
    possessive = r"(?:của|cua)"
    entity_label = r"(?:lớp|lop|giáo\s*viên|giao\s*vien|gv)"
    text = re.sub(
        rf"^\s*{timetable_label}\s*(?:{possessive}\s*)?(?:{entity_label})?\s*[:\-]?\s*",
        "",
        text,
        flags=re.IGNORECASE,
    ).strip()
    text = re.sub(
        rf"^\s*{entity_label}\s*[:\-]?\s*", "", text, flags=re.IGNORECASE
    ).strip()
    return text or original


def _standalone_class_token(value: str) -> str:
    text = _cell_text(value)
    if not text:
        return ""
    patterns = [
        r"(?i)(?:\blop\s*)?(\d{1,2}\s*[A-Z]{1,3}\s*\d{0,2})\b",
        r"(?i)(?:\blop\s*)?(\d{1,2}\s*[/\-]\s*\d{1,2})\b",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return re.sub(r"\s+", "", match.group(1)).upper()
    normalized = normalize_text(text)
    match = re.fullmatch(r"(?:lop\s*)?(\d{1,2})\s*([a-z]{1,3})\s*(\d{0,2})", normalized)
    if match:
        return f"{match.group(1)}{match.group(2).upper()}{match.group(3)}"
    return ""


def _standalone_looks_like_class(value: str) -> bool:
    return bool(_standalone_class_token(value))


def _standalone_without_class_token(value: str, class_name: str = "") -> str:
    """Remove one class token from a mixed cell while preserving Vietnamese text."""
    text = _cell_text(value)
    target = normalize_text(class_name or _standalone_class_token(text))
    if not text or not target:
        return text
    patterns = (
        r"(?i)(?:\b(?:lớp|lop)\s*)?(\d{1,2}\s*[A-Z]{1,3}\s*\d{0,2})\b",
        r"(?i)(?:\b(?:lớp|lop)\s*)?(\d{1,2}\s*[/\-]\s*\d{1,2})\b",
    )
    for pattern in patterns:
        match = re.search(pattern, text)
        if not match or normalize_text(re.sub(r"\s+", "", match.group(1))).replace(
            " ", ""
        ) != target.replace(" ", ""):
            continue
        left = text[: match.start()].strip(" -–—|;/:,·•")
        right = text[match.end() :].strip(" -–—|;/:,·•")
        return " ".join(part for part in (left, right) if part).strip()
    return text


def _standalone_subject_from_text(value: str) -> str:
    subject, _teacher = _standalone_subject_prefix(value)
    if subject:
        return subject
    text = _cell_text(value)
    if not text:
        return ""
    explicit = re.search(r"(?i)\b(?:môn|mon|subject)\s*[:\-]\s*([^|;/]+)", text)
    if explicit:
        explicit_text = explicit.group(1).strip()
        explicit_subject, _remainder = _standalone_subject_prefix(explicit_text)
        if explicit_subject:
            return explicit_subject
        explicit_norm = normalize_text(explicit_text)
        if explicit_norm in _STANDALONE_SUBJECT_HINTS:
            return _STANDALONE_SUBJECT_HINTS[explicit_norm]
        if explicit_text:
            return explicit_text

    # Prefer structured pieces and prefixes. Searching every token in the whole cell
    # can mistake a teacher's middle name (for example "Nguyễn Văn An") for môn Văn.
    candidates = _standalone_split_parts(text)
    class_name = _standalone_class_token(text)
    full_norm = normalize_text(text)
    class_norm = normalize_text(class_name)
    normalized_candidates = [normalize_text(part) for part in candidates]
    if class_norm and full_norm.startswith(f"{class_norm} "):
        normalized_candidates.append(full_norm[len(class_norm) :].strip())

    for part in candidates:
        part_subject, _remainder = _standalone_subject_prefix(part)
        if part_subject:
            return part_subject

    for norm in normalized_candidates:
        if not norm:
            continue
        for hint in sorted(_STANDALONE_SUBJECT_HINTS, key=len, reverse=True):
            if norm == hint or norm.startswith(f"{hint} "):
                return _STANDALONE_SUBJECT_HINTS[hint]
    return ""


def _standalone_split_parts(value: str) -> list[str]:
    text = _cell_text(value)
    if not text:
        return []
    parts = [
        part.strip()
        for part in re.split(r"\s*(?:\n|\r|\||;|·|•|\s[-–—]\s|\s/\s)\s*", text)
        if part.strip()
    ]
    return parts or [text]


def _standalone_teacher_from_text(value: str, *, exclude: Iterable[str] = ()) -> str:
    text = _cell_text(value)
    if not text:
        return ""
    explicit = re.search(
        r"(?i)\b(?:giáo\s*viên|giao\s*vien|gv)\s*[:\-]\s*([^|;/]+)", text
    )
    if explicit:
        return explicit.group(1).strip()
    excluded = {normalize_text(item) for item in exclude if item}
    _subject, remainder = _standalone_subject_prefix(text)
    if (
        remainder
        and normalize_text(remainder) not in excluded
        and not _standalone_looks_like_class(remainder)
    ):
        return remainder
    parts = _standalone_split_parts(text)
    for part in reversed(parts):
        norm = normalize_text(part)
        if not norm or norm in excluded or _standalone_looks_like_class(part):
            continue
        # Only reject a part when the part itself reads like a subject. Looking
        # for subject words anywhere would drop valid names such as Nguyen Van An.
        if _standalone_subject_from_text(part):
            continue
        # Ten giao vien thuong co it nhat 2 tu; chap nhan ma GV viet hoa ngan neu co tien to GV.
        if len(norm.split()) >= 2:
            return part.strip()
    return ""


def _standalone_parse_cell(
    value: str,
    *,
    fixed_class: str = "",
    fixed_teacher: str = "",
) -> tuple[str, str, str]:
    text = _cell_text(value)
    class_name = fixed_class or _standalone_class_token(text)
    # Mixed cells such as "10A1 Toán Nguyễn Văn An" must be parsed from the
    # content after the class token. Otherwise the subject prefix can prevent
    # the remaining teacher name from being recognized.
    semantic_text = _standalone_without_class_token(text, class_name)
    subject_name = _standalone_subject_from_text(semantic_text)
    teacher_name = fixed_teacher
    if not teacher_name:
        teacher_name = _standalone_teacher_from_text(
            semantic_text, exclude=(class_name, subject_name)
        )
    parts = _standalone_split_parts(semantic_text)
    if not subject_name:
        for part in parts:
            if class_name and normalize_text(part) == normalize_text(class_name):
                continue
            if teacher_name and normalize_text(part) == normalize_text(teacher_name):
                continue
            if _standalone_looks_like_class(part):
                continue
            norm = normalize_text(part)
            if norm and not norm.startswith(("gv ", "giao vien ")):
                subject_name = part.strip()
                break
    # Khong tao entity tam tai buoc parse. De trong de tang phan analyze co the
    # canh bao ro rang va khong dem "chua xac dinh" nhu mon/giao vien that.
    return class_name.strip(), subject_name.strip(), teacher_name.strip()


_STANDALONE_OPTIONAL_TEACHER_SUBJECTS = {
    "chao co",
    "sinh hoat",
    "trai nghiem",
    "hdtn",
    "hdtnhn",
    "tnhn",
}


def _standalone_teacher_is_optional(subject_name: str) -> bool:
    return normalize_text(subject_name) in _STANDALONE_OPTIONAL_TEACHER_SUBJECTS


def _standalone_day_index(value: str) -> int | None:
    text = normalize_text(value)
    if not text:
        return None
    exact = {
        "cn": 6,
        "chu nhat": 6,
        "sunday": 6,
        "sun": 6,
        "monday": 0,
        "mon": 0,
        "tuesday": 1,
        "tue": 1,
        "wednesday": 2,
        "wed": 2,
        "thursday": 3,
        "friday": 4,
        "fri": 4,
        "saturday": 5,
        "sat": 5,
    }
    if text in exact:
        return exact[text]
    match = re.fullmatch(r"(?:thu\s*)?([2-8])", text)
    if match:
        number = int(match.group(1))
        return 6 if number == 8 else number - 2
    return None


def _standalone_heading_from_context(
    title: str, rows: list[list[str]], header_index: int
) -> str:
    """Pick the class/teacher heading from a sheet title or rows above the grid header."""
    generic = {"", "du lieu", "data", "thoi khoa bieu", "tkb", "sheet"}

    def is_generic_heading(normalized: str) -> bool:
        return (
            normalized in generic
            or re.fullmatch(r"sheet\s*\d+", normalized) is not None
        )

    title_heading = _standalone_clean_entity_heading(title)
    title_norm = normalize_text(title_heading)

    # A meaningful sheet/table title is the safest source. Excel's default
    # Sheet1/Sheet2/... names are metadata, not class/teacher identities.
    if not is_generic_heading(title_norm):
        return title_heading

    # Excel files often keep a generic sheet name and put the real heading in a
    # merged cell directly above the weekday row. Scan bottom-up so the nearest
    # heading wins over school names or other document banners.
    fallback = ""
    for row in reversed(rows[:header_index]):
        for value in row:
            raw = _cell_text(value)
            if not raw:
                continue
            cleaned = _standalone_clean_entity_heading(raw)
            cleaned_norm = normalize_text(cleaned)
            if not cleaned_norm or is_generic_heading(cleaned_norm):
                continue
            raw_norm = normalize_text(raw)
            has_entity_label = (
                "thoi khoa bieu" in raw_norm
                or raw_norm.startswith("giao vien ")
                or raw_norm.startswith("gv ")
                or raw_norm.startswith("lop ")
            )
            if has_entity_label or _standalone_looks_like_class(cleaned):
                return cleaned
            if not fallback and len(cleaned.split()) <= 8:
                fallback = cleaned
    return fallback or title_heading


def _standalone_find_grid_header(
    rows: list[list[str]],
) -> tuple[int, list[tuple[int, str]], int, int | None] | None:
    for header_index, row in enumerate(rows[:30]):
        day_cols = [
            (col, _cell_text(value))
            for col, value in enumerate(row)
            if _standalone_day_index(_cell_text(value)) is not None
        ]
        distinct_days = {_standalone_day_index(day_text) for _col, day_text in day_cols}
        if len(day_cols) < 2 or len(distinct_days) < 2:
            continue
        period_header_col = next(
            (
                col
                for col, value in enumerate(row)
                if _header_kind(_cell_text(value)) == "period"
            ),
            None,
        )
        # Numeric long/wide rows such as "2, 4, ..." must not be mistaken for a
        # grid header merely because both numbers can also mean weekdays.  A
        # two-day grid therefore needs an explicit Tiết/Period header; header
        # rows without one are accepted only when at least three distinct days
        # make the grid shape unambiguous.
        if period_header_col is None and len(distinct_days) < 3:
            continue
        period_col = period_header_col if period_header_col is not None else 0
        session_col = next(
            (
                col
                for col, value in enumerate(row)
                if _header_kind(_cell_text(value)) == "session"
            ),
            None,
        )
        return header_index, day_cols, period_col, session_col
    return None


def _standalone_parse_grid(
    title: str, rows: list[list[str]]
) -> tuple[list[RawLesson], str]:
    header = _standalone_find_grid_header(rows)
    if header is None:
        return [], ""
    header_index, day_cols, period_col, session_col = header
    heading = _standalone_heading_from_context(title, rows, header_index)
    sample_cells: list[str] = []
    for values in rows[header_index + 1 : header_index + 18]:
        for col, _day in day_cols:
            if col < len(values) and _cell_text(values[col]):
                sample_cells.append(_cell_text(values[col]))
    class_hits = sum(1 for cell in sample_cells if _standalone_class_token(cell))
    is_class_sheet = _standalone_looks_like_class(heading)
    if not is_class_sheet and sample_cells:
        is_teacher_sheet = class_hits >= max(1, len(sample_cells) // 3)
    else:
        is_teacher_sheet = False
    if not is_class_sheet and not is_teacher_sheet:
        # Ten sheet ngan khong giong tieu de chung thuong la ten lop; neu la ten nguoi
        # ma o co lop thi nhanh teacher o tren da bat duoc.
        is_class_sheet = bool(
            heading
            and normalize_text(heading)
            not in {"du lieu", "thoi khoa bieu", "tkb", "sheet", "sheet1"}
        )
    if not is_class_sheet and not is_teacher_sheet:
        return [], ""

    parsed: list[RawLesson] = []
    last_session = ""
    implicit_sessions = _implicit_session_hints(
        rows,
        start_index=header_index + 1,
        period_col=period_col,
        session_col=session_col,
    )
    for row_index, values in enumerate(
        rows[header_index + 1 :], start=header_index + 1
    ):
        row_no = row_index + 1
        period = values[period_col] if period_col < len(values) else ""
        if not period or _period_number(period) is None:
            continue
        session_value = (
            values[session_col]
            if session_col is not None and session_col < len(values)
            else ""
        )
        if session_value:
            last_session = session_value
        session = session_value or implicit_sessions.get(row_index, "") or last_session
        for col, day_text in day_cols:
            lesson = values[col] if col < len(values) else ""
            if not lesson or normalize_text(lesson) in {
                "x",
                "trong",
                "nghi",
                "off",
                "none",
                "na",
            }:
                continue
            if is_teacher_sheet:
                class_name, subject_name, teacher_name = _standalone_parse_cell(
                    lesson, fixed_teacher=heading
                )
                if not class_name:
                    continue
                parsed.append(
                    RawLesson(
                        day_text,
                        session,
                        period,
                        class_name,
                        lesson_text=lesson,
                        subject_text=subject_name,
                        teacher_text=teacher_name,
                        source=f"{title} · dong {row_no}",
                        origin="teacher",
                    )
                )
            else:
                class_name, subject_name, teacher_name = _standalone_parse_cell(
                    lesson, fixed_class=heading
                )
                parsed.append(
                    RawLesson(
                        day_text,
                        session,
                        period,
                        class_name or heading,
                        lesson_text=lesson,
                        subject_text=subject_name,
                        teacher_text=teacher_name,
                        source=f"{title} · dong {row_no}",
                        origin="class",
                    )
                )
    if not parsed:
        return [], ""
    return parsed, "Theo giao vien" if is_teacher_sheet else "Theo lop/hoc sinh"


def _standalone_parse_wide(
    title: str, rows: list[list[str]]
) -> tuple[list[RawLesson], str]:
    best = None
    for header_index, row in enumerate(rows[:30]):
        kinds: dict[str, int] = {}
        for col, value in enumerate(row):
            kind = _header_kind(_cell_text(value))
            if kind in {"day", "session", "period"} and kind not in kinds:
                kinds[kind] = col
        if "day" not in kinds or "period" not in kinds:
            continue
        entity_cols = [
            (col, _cell_text(value))
            for col, value in enumerate(row)
            if col not in kinds.values()
            and _cell_text(value)
            and not _is_wide_table_auxiliary_header(_cell_text(value))
        ]
        if not entity_cols:
            continue
        score = len(entity_cols)
        if best is None or score > best[0]:
            best = (score, header_index, kinds, entity_cols)
    if best is None:
        return [], ""
    _score, header_index, kinds, entity_cols = best
    class_header_hits = sum(
        1 for _col, header in entity_cols if _standalone_looks_like_class(header)
    )
    sampled = []
    for values in rows[header_index + 1 : header_index + 18]:
        for col, _header in entity_cols:
            if col < len(values) and _cell_text(values[col]):
                sampled.append(_cell_text(values[col]))
    cell_class_hits = sum(1 for cell in sampled if _standalone_class_token(cell))
    teacher_wide = (
        class_header_hits == 0 and cell_class_hits >= max(1, len(sampled) // 3)
        if sampled
        else False
    )

    parsed: list[RawLesson] = []
    last_day = ""
    last_session = ""
    implicit_sessions = _implicit_session_hints(
        rows,
        start_index=header_index + 1,
        period_col=kinds["period"],
        session_col=kinds.get("session"),
        day_col=kinds["day"],
    )
    for row_index, values in enumerate(
        rows[header_index + 1 :], start=header_index + 1
    ):
        row_no = row_index + 1
        get = lambda col: values[col] if col is not None and col < len(values) else ""
        day_value = get(kinds.get("day"))
        session_value = get(kinds.get("session"))
        period = get(kinds.get("period"))
        if day_value:
            last_day = day_value
        if session_value:
            last_session = session_value
        day = day_value or last_day
        session = session_value or implicit_sessions.get(row_index, "") or last_session
        if not day or not period or _period_number(period) is None:
            continue
        for col, header in entity_cols:
            lesson = get(col)
            if not lesson or normalize_text(lesson) in {
                "x",
                "trong",
                "nghi",
                "off",
                "none",
                "na",
            }:
                continue
            if teacher_wide:
                class_name, subject_name, teacher_name = _standalone_parse_cell(
                    lesson,
                    fixed_teacher=_standalone_clean_entity_heading(header),
                )
                if not class_name:
                    continue
                parsed.append(
                    RawLesson(
                        day,
                        session,
                        period,
                        class_name,
                        lesson_text=lesson,
                        subject_text=subject_name,
                        teacher_text=teacher_name,
                        source=f"{title} · dong {row_no}",
                        origin="teacher",
                    )
                )
            else:
                class_name, subject_name, teacher_name = _standalone_parse_cell(
                    lesson, fixed_class=header
                )
                parsed.append(
                    RawLesson(
                        day,
                        session,
                        period,
                        class_name or header,
                        lesson_text=lesson,
                        subject_text=subject_name,
                        teacher_text=teacher_name,
                        source=f"{title} · dong {row_no}",
                        origin="class",
                    )
                )
    if not parsed:
        return [], ""
    return (
        parsed,
        "Bang tong hop theo giao vien" if teacher_wide else "Bang tong hop theo lop",
    )


def parse_tables_standalone(
    tables: list[tuple[str, list[list[str]]]],
) -> tuple[list[RawLesson], list[str], dict[str, Any]]:
    all_rows: list[RawLesson] = []
    warnings: list[str] = []
    layouts: set[str] = set()
    for title, raw_rows in tables:
        rows = _trim_matrix(raw_rows)
        if not rows:
            continue
        parsed = _parse_long_table(title, rows)
        layout = "Bang tong hop" if parsed else ""
        if parsed:
            enriched = []
            for item in parsed:
                class_name, subject_name, teacher_name = _standalone_parse_cell(
                    item.lesson_text or item.subject_text or item.teacher_text,
                    fixed_class=item.class_text,
                    fixed_teacher=item.teacher_text,
                )
                item.class_text = class_name or item.class_text
                item.subject_text = item.subject_text or subject_name
                item.teacher_text = item.teacher_text or teacher_name
                enriched.append(item)
            parsed = enriched
        if not parsed:
            parsed, layout = _standalone_parse_grid(title, rows)
        if not parsed:
            parsed, layout = _standalone_parse_wide(title, rows)
        if parsed:
            all_rows.extend(parsed)
            layouts.add(layout)
        elif any(
            _header_kind(value) in {"day", "period"}
            for row in rows[:30]
            for value in row
        ):
            warnings.append(
                f'Khong nhan dien duoc cau truc thoi khoa bieu trong "{title}".'
            )
    return all_rows, warnings, {"layouts": sorted(layouts)}


__all__ = [name for name in globals() if not name.startswith("__")]
